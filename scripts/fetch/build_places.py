"""COLM 官方「舊金山在地資訊」試算表 + Google My Maps → data-src/places.json

兩份來源互補：
  * 試算表（Dan Jurafsky 整理）：分類、價位、描述、每個地點的 Google Maps 連結（cid）
  * My Maps 的 KML 匯出：同一批地點的經緯度（試算表沒有座標）

以試算表為主（描述較新、有價位），用「分類 + 名稱」對上 KML 取座標。
同一個地點出現在兩個分類（Taishan、Maison Nico）會合併成一筆、cats 有兩個。

Usage（在 repo 根目錄）:
    uv run --with openpyxl --with certifi python scripts/fetch/build_places.py
    uv run --with openpyxl python scripts/fetch/build_places.py --xlsx path/to/local.xlsx --kml path/to/map.kml
"""
from __future__ import annotations

import argparse
import json
import re
import ssl
import unicodedata
import urllib.request
import xml.etree.ElementTree as ET
from pathlib import Path

import openpyxl

try:                                    # uv 裝的 Python 不讀 macOS 鑰匙圈，用 certifi 的 CA
    import certifi
    SSL_CTX = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    SSL_CTX = ssl.create_default_context()

ROOT = Path(__file__).resolve().parents[2]
SHEET_ID = "1xy8HOYJOyWzFqnsmr06t-NcVoLKDuwqb1A1Es6iaY3I"
MYMAPS_MID = "1U_Po7vWtGCGrEnHJ3LWl_hbvMd0TWEo"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit"
XLSX_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=xlsx"
KML_URL = f"https://www.google.com/maps/d/kml?mid={MYMAPS_MID}&forcekml=1"
MYMAPS_URL = f"https://www.google.com/maps/d/viewer?mid={MYMAPS_MID}"

# 分頁名稱（會被 Google 截到 31 字）的前綴 → 分類 id。KML 資料夾名稱大小寫略有不同，一樣用前綴對。
CATS = [
    ("food",      "restaurants nearby",          "Restaurants Nearby"),
    ("latenight", "late night restaurants",      "Late Night Restaurants Nearby"),
    ("cafe",      "boba, bakeries and cafes",    "Boba, Bakeries and Cafes Nearby"),
    ("bars",      "bars nearby",                 "Bars Nearby"),
    ("culture",   "museums, art, tourist",       "Museums, Art, Tourist Attractions"),
    ("parks",     "parks nearby",                "Parks Nearby"),
    ("further",   "more distant parks",          "More Distant Parks, Museums, Trails"),
]
KML_PREFIX = {"nearby museums": "culture"}      # KML 的資料夾名和分頁名不一樣的只有這個


def cat_of(title: str) -> str | None:
    t = title.lower().strip()
    for prefix, cid in KML_PREFIX.items():
        if t.startswith(prefix):
            return cid
    for cid, prefix, _ in CATS:
        if t.startswith(prefix):
            return cid
    return None


def norm(name: str) -> str:
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", s)


def slug(name: str) -> str:
    s = re.sub(r"['’]", "", name)                       # Tony's → tonys，不要 tony-s
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def fetch(url: str, dest: Path) -> Path:
    req = urllib.request.Request(url, headers={"User-Agent": "colm2026-site/1.0"})
    with urllib.request.urlopen(req, timeout=60, context=SSL_CTX) as r:
        dest.write_bytes(r.read())
    return dest


def clean(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip()


def read_sheet(xlsx: Path) -> list[dict]:
    """→ rows: {cat, name, price, desc, url}。Intro 分頁是散文，手動整理在 data-src/map-text.json。"""
    wb = openpyxl.load_workbook(xlsx)
    rows = []
    for ws in wb.worksheets:
        if ws.title.lower().startswith("intro"):
            continue
        cid = cat_of(ws.title)
        if not cid:
            print(f"  !! 不認得的分頁：{ws.title}")
            continue
        header = [clean(c.value).lower() for c in ws[1]]
        for r in ws.iter_rows(min_row=2):
            cells = dict(zip(header, r))
            name_cell = cells.get("name")
            if not name_cell or not name_cell.value:
                continue
            price = cells.get("price")
            desc = cells.get("description")
            rows.append({
                "cat": cid,
                "name": clean(name_cell.value),
                "price": clean(price.value) if price is not None and price.value else None,
                "desc": clean(desc.value) if desc is not None and desc.value else None,
                "url": name_cell.hyperlink.target if name_cell.hyperlink else None,
            })
    return rows


def read_kml(kml: Path) -> dict[tuple[str, str], dict]:
    """→ {(cat, normname): {lat, lng, desc}}"""
    ns = {"k": "http://www.opengis.net/kml/2.2"}
    root = ET.parse(kml).getroot()
    out = {}
    for folder in root.iter("{http://www.opengis.net/kml/2.2}Folder"):
        cid = cat_of(folder.find("k:name", ns).text or "")
        for pm in folder.findall("k:Placemark", ns):
            name = clean(pm.find("k:name", ns).text)
            coords = pm.find(".//k:coordinates", ns)
            if coords is None:
                continue
            lng, lat, *_ = (float(x) for x in coords.text.strip().split(","))
            d = pm.find("k:description", ns)
            out[(cid, norm(name))] = {"lat": round(lat, 6), "lng": round(lng, 6), "desc": clean(d.text) if d is not None else None}
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", help="本機的試算表 .xlsx（不給就從 Google 下載）")
    ap.add_argument("--kml", help="本機的 My Maps .kml（不給就從 Google 下載）")
    ap.add_argument("--out", default=str(ROOT / "data-src/places.json"))
    a = ap.parse_args()

    work = ROOT / "tmp/places-raw"            # tmp/ 不進版控
    work.mkdir(parents=True, exist_ok=True)
    xlsx = Path(a.xlsx) if a.xlsx else fetch(XLSX_URL, work / "sheet.xlsx")
    kml = Path(a.kml) if a.kml else fetch(KML_URL, work / "map.kml")

    rows = read_sheet(xlsx)
    geo = read_kml(kml)

    places: dict[str, dict] = {}             # key = Google Maps 連結（同一地點跨分類時相同）
    missing = []
    for r in rows:
        g = geo.get((r["cat"], norm(r["name"])))
        if not g:
            missing.append(f"{r['cat']}: {r['name']}")
            continue
        key = r["url"] or norm(r["name"])
        if key in places:
            p = places[key]
            p["cats"].append(r["cat"])
            p["price"] = p["price"] or r["price"]
            p["desc"] = p["desc"] or r["desc"] or g["desc"]
            continue
        places[key] = {
            "id": slug(r["name"]),
            "name": r["name"],
            "cats": [r["cat"]],
            "price": r["price"],
            "desc": r["desc"] or g["desc"],
            "lat": g["lat"], "lng": g["lng"],
            "gmaps": r["url"],
        }

    data = {
        "source": {"sheet": SHEET_URL, "mymaps": MYMAPS_URL,
                   "note": "COLM 2026 SF local information, curated by Dan Jurafsky for COLM 2026."},
        "categories": [{"id": cid, "en": label} for cid, _, label in CATS],
        "places": list(places.values()),
    }
    Path(a.out).write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    per_cat = {cid: sum(1 for p in data["places"] if cid in p["cats"]) for cid, _, _ in CATS}
    print(f"places.json  {len(data['places'])} 個地點  {per_cat}")
    for m in missing:
        print(f"  !! KML 找不到座標：{m}")


if __name__ == "__main__":
    main()
